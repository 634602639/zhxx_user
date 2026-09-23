from __future__ import annotations

import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


SOURCE_DIR = Path(r"C:\Users\21943\Desktop\截图")
OUTPUT_PATH = SOURCE_DIR / "简述.docx"
TEMP_PATH = Path(r"E:\zhxx_user\.codex-docx-work\简述.tmp.docx")

ITEMS = [
    (
        "01 登录与身份认证",
        SOURCE_DIR / "登录.png",
        "登录界面",
        "登录界面是综合信息服务中心的统一身份入口，采用居中卡片式布局，提供用户名、密码输入和登录操作。系统在提交后校验账号与密码，并通过会话保存登录状态；未登录或会话失效时会自动回到该页面。简洁的绿色视觉与主系统保持一致，既突出操作入口，也减少无关信息干扰，为后续数据采集、分析、报告生成及管理功能提供安全、清晰的访问起点。",
    ),
    (
        "02 多源数据采集与管理",
        SOURCE_DIR / "多源数据采集与管理.png",
        "多源数据采集与管理界面",
        "该界面承担多源数据接入、整理与入库管理。上方可按单位维护基地接口，设置时间范围后批量拉取指标；远程采集配置支持新增、编辑、删除及按顺序执行。采集结果先进入解析预览，确认后再写入标签数据区，并按数据日期聚合展示。用户可查看来源、类别、采集时间和业务时间，也能修改或删除记录，从而形成可追溯、可复用的数据底座，为后续分析与报告填报提供统一数据来源。",
    ),
    (
        "03 值勤数据分析与展示",
        SOURCE_DIR / "值勤数据分析与展示.png",
        "值勤数据分析与展示界面",
        "值勤数据分析界面将已入库指标转化为可视化态势。顶部支持按时间范围、单位和图表类型筛选，关键指标卡集中呈现节点在线率、设备完好率、入网用户和在线节点数量。下方折线图对运行百分率、业务量及节点变化进行趋势对比，便于快速发现波动和差异。当指定范围缺少入库数据时，系统可展示预置指标作为演示，并在界面中明确提示，帮助用户完成态势研判和报告数据核对。",
    ),
    (
        "04 值勤报告辅助生成",
        SOURCE_DIR / "值勤辅助报告生成.png",
        "值勤报告辅助生成界面",
        "报告辅助生成界面覆盖模板上传、占位填报和草稿生成三个环节。用户可上传 Word 或 PowerPoint 模板，系统自动提取关键词并建立模板清单；选择时间、单位和指标后，可将已入库数据映射到模板占位符，也可启用大模型辅助填写。生成结果保存为合成草稿，支持润色纠错、下载预览、正式保存和删除，并保留处理状态，减少重复复制工作，提高报告编制的规范性和效率。",
    ),
    (
        "05 值勤报告在线管理",
        SOURCE_DIR / "值勤报告在线管理.png",
        "值勤报告在线管理界面",
        "值勤报告在线管理界面用于集中维护正式归档的报告。用户可按标题关键字、起止日期和文稿类型组合查询，列表显示报告编号、标题、类型、日期及可用操作。除从本地新增报告外，还可对现有报告进行编辑、导出 Word、生成 PDF 或删除。该模块承接辅助生成环节的正式保存结果，实现报告从形成、检索、修订到分发的闭环，便于定位历史材料并保持文档版本有序。",
    ),
    (
        "06 操作日志",
        SOURCE_DIR / "操作日志.png",
        "操作日志界面",
        "操作日志界面记录系统关键行为，为运行追踪和责任审计提供依据。日志按时间倒序展示编号、功能模块、操作类型、详细参数、执行结果和发生时间，并支持按模块、操作或详情关键字检索。数据采集、分析查询、报告管理、用户登录退出等行为均可形成记录。管理员可据此核查操作路径、定位异常和统计使用情况，也为数据变更及报告流转提供完整的追溯线索。",
    ),
]


def set_run_font(run, east_asia="Microsoft YaHei", latin="Arial", size=None, bold=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, text, fld_end])
    set_run_font(run, size=8.5, color=(110, 118, 130))


def set_picture_alt_text(run, title, description):
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        return
    doc_pr = drawing.find(".//" + qn("wp:docPr"))
    if doc_pr is not None:
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def remove_paragraph_borders(paragraph_or_style):
    if hasattr(paragraph_or_style, "_p"):
        ppr = paragraph_or_style._p.get_or_add_pPr()
    else:
        ppr = paragraph_or_style._element.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is not None:
        ppr.remove(borders)


def add_section_page(doc, index, heading, image_path, caption, description):
    if index > 0:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.keep_with_next = True
    h.add_run(heading)

    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(1)
    image_p.paragraph_format.space_after = Pt(3)
    image_p.paragraph_format.keep_with_next = True
    image_run = image_p.add_run()
    image_run.add_picture(str(image_path), width=Mm(242))
    set_picture_alt_text(image_run, caption, f"{caption}截图")

    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap_run = cap.add_run(f"图{index}  {caption}")
    set_run_font(cap_run, size=9, color=(105, 114, 126))

    body = doc.add_paragraph(style="Brief Description")
    body.paragraph_format.keep_together = True
    lead = body.add_run("简述  ")
    set_run_font(lead, size=10.5, bold=True, color=(36, 132, 84))
    run = body.add_run(description)
    set_run_font(run, size=10.5, color=(31, 41, 55))


def build_document():
    missing = [str(path) for _, path, _, _ in ITEMS if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing screenshots: " + ", ".join(missing))

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(7)

    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(5)

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)
    remove_paragraph_borders(title_style)

    heading = styles["Heading 1"]
    heading.font.name = "Arial"
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    heading.font.size = Pt(18)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True

    caption_style = styles["Caption"]
    caption_style.font.name = "Arial"
    caption_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption_style.font.size = Pt(9)
    caption_style.font.italic = False
    caption_style.font.color.rgb = RGBColor(105, 114, 126)
    caption_style.paragraph_format.space_before = Pt(0)
    caption_style.paragraph_format.space_after = Pt(5)

    if "Brief Description" not in [s.name for s in styles]:
        brief_style = styles.add_style("Brief Description", WD_STYLE_TYPE.PARAGRAPH)
    else:
        brief_style = styles["Brief Description"]
    brief_style.base_style = normal
    brief_style.font.name = "Arial"
    brief_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    brief_style.font.size = Pt(10.5)
    brief_style.font.color.rgb = RGBColor(31, 41, 55)
    brief_style.paragraph_format.line_spacing = 1.3
    brief_style.paragraph_format.space_before = Pt(4)
    brief_style.paragraph_format.space_after = Pt(0)
    brief_style.paragraph_format.first_line_indent = Mm(0)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    footer_run = footer_p.add_run("综合信息服务中心系统界面简述    第 ")
    set_run_font(footer_run, size=8.5, color=(110, 118, 130))
    add_page_field(footer_p)
    footer_tail = footer_p.add_run(" 页")
    set_run_font(footer_tail, size=8.5, color=(110, 118, 130))

    doc.core_properties.title = "综合信息服务中心系统界面简述"
    doc.core_properties.subject = "系统主要界面截图及功能说明"
    doc.core_properties.keywords = "综合信息服务中心, 数据采集, 值勤分析, 报告生成, 操作日志"

    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    title.add_run("综合信息服务中心系统界面简述")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    sub_run = subtitle.add_run("界面截图及功能说明")
    set_run_font(sub_run, size=14, color=(72, 84, 99))

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.left_indent = Mm(30)
    intro.paragraph_format.right_indent = Mm(30)
    intro.paragraph_format.first_line_indent = Mm(8)
    intro.paragraph_format.line_spacing = 1.55
    intro.paragraph_format.space_after = Pt(18)
    intro_text = (
        "本文依据系统实际界面截图和项目功能，对登录、数据采集、值勤分析、报告辅助生成、报告在线管理与操作日志六个主要界面进行说明。"
        "各页按业务流程排列，便于快速了解系统从身份认证、数据汇聚、态势研判到报告形成和审计追溯的完整链路。"
    )
    intro_run = intro.add_run(intro_text)
    set_run_font(intro_run, size=12, color=(31, 41, 55))

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope_run = scope.add_run("登录    数据采集    数据分析    报告生成    报告管理    操作日志")
    set_run_font(scope_run, size=11, bold=True, color=(36, 132, 84))

    for idx, (heading_text, image_path, caption, description) in enumerate(ITEMS, start=1):
        add_section_page(doc, idx, heading_text, image_path, caption, description)

    TEMP_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TEMP_PATH)
    shutil.copy2(TEMP_PATH, OUTPUT_PATH)
    TEMP_PATH.unlink()

    print(f"Created: {OUTPUT_PATH}")
    for heading_text, _, _, description in ITEMS:
        print(f"{heading_text}: {len(description)} characters")


if __name__ == "__main__":
    build_document()
