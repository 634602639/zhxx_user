"""值勤报告输出：Word(.docx) / PDF（按正文文字生成）

- Word：按正文生成 .docx；存在 file_blob 的报告由路由直接回传原始附件。
- PDF：用 reportlab 把标题 + 正文文字渲染为 PDF（自带中文 CID 字体，
  无需 Word/LibreOffice，也不还原 .docx 原排版）。
"""
import os

from docx import Document
from docx.shared import Pt


def export_word(report, out_dir: str) -> str:
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"report_{report.id}.docx")
    doc = Document()
    doc.add_heading(report.title or "值勤报告", level=1)
    for line in (report.content or "").split("\n"):
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.size = Pt(11)
    doc.save(path)
    return path


# 中文字体只需注册一次（STSong-Light 为 reportlab 自带的 Adobe 亚洲 CID 字体）
_CJK_FONT = "STSong-Light"
_cjk_font_ready = False


def _ensure_cjk_font():
    global _cjk_font_ready
    if _cjk_font_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont(_CJK_FONT))
    _cjk_font_ready = True


def _pdf_escape(text: str) -> str:
    """reportlab Paragraph 按简易 XML 解析，需转义 & < >。"""
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def export_pdf(report, out_dir: str) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    _ensure_cjk_font()
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"report_{report.id}.pdf")

    title_style = ParagraphStyle(
        "CJKTitle", fontName=_CJK_FONT, fontSize=18, leading=26,
        spaceAfter=12, alignment=1,  # 居中
    )
    body_style = ParagraphStyle(
        "CJKBody", fontName=_CJK_FONT, fontSize=11, leading=18,
    )

    doc = SimpleDocTemplate(
        path, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=(report.title or "值勤报告"),
    )
    story = [Paragraph(_pdf_escape(report.title or "值勤报告"), title_style), Spacer(1, 6)]
    for line in (report.content or "").split("\n"):
        # 空行用占位空格保持段间距，避免 reportlab 忽略空 Paragraph
        story.append(Paragraph(_pdf_escape(line) or "&nbsp;", body_style))
    doc.build(story)
    return path
