# -*- coding: utf-8 -*-
"""把 7.1.4.1.1 多源数据采集与管理 / 数据提取 的测试用例表补充进考核大纲。
以 7.1.1（排班查询 TABLE#13）为模板克隆，保持完全一致的表格格式。"""
import copy
import docx
from docx.oxml.ns import qn

SRC = "系统出厂（所）考核大纲模版.docx"

CASE_ID = "数据提取/GN_YWBG_ZHXXFWZXYY_DYSJCJYGL_SJTQ"
CASE_DESC = (
    "从运维管理系统等外部业务系统采集值勤报告所需的原始数据，验证支持新增/编辑/删除远程采集配置、"
    "按配置采集或顺序采集全部、解析预览（不入库）、选中采集原文暂存为标签、入库及按天聚合查看的"
    "多源数据采集与管理全流程"
)

# 每个步骤 = [前提和约束, 输入, 目的和动作, 预期结果, 评估准则, 备注]
# 输入/预期等可为 list[str] 表示多行
STEPS = [
    ["已在浏览器中登录进入综合信息服务中心应用软件", "",
     "点击左侧功能菜单中的“多源数据采集”菜单",
     "进入“多源数据采集与管理”界面，界面自上而下依次显示远程采集配置、解析预览、标签暂存区、已入库数据等区域。",
     "与预期结果一致", "-"],

    ["", "",
     "点击界面右上方的“新增配置”按钮",
     "弹出“新增采集配置”面板，包含配置名称、接口URL、默认来源source、默认类别category等输入项。",
     "与预期结果一致", ""],

    ["接口URL可正常访问并返回JSON数据，接口返回示例见下方前提数据",
     ["配置名称：运维告警接口、", "接口URL：http://192.168.11.87:8080/api/ops/alarms、",
      "默认来源：运维管理系统、", "默认类别：告警"],
     "在“新增采集配置”面板依次输入配置名称、接口URL、默认来源、默认类别，点击“保存配置”按钮。",
     "弹出成功提示“已新增”，面板关闭，远程采集配置列表新增一条配置，信息与输入一致，序号自动生成并排在末尾。",
     "与预期结果一致", ""],

    ["已存在步骤3新增的采集配置", "配置名称：运维告警接口A",
     "点击该配置后面的“编辑”按钮，修改配置名称，点击“保存配置”按钮。",
     "弹出成功提示“已更新”，远程采集配置列表显示修改后的数据，配置名称与输入一致。",
     "与预期结果一致", ""],

    ["", "",
     "点击该配置后面的“采集”按钮",
     ["弹出成功提示“解析成功，共 N 条（未入库）”，解析预览区显示采集到的原始数据列表（不入库），",
      "列含序号、所属配置、方式、来源、类别、标题、时间，记录条数及内容与接口返回数据一致。"],
     "与预期结果一致", "数据提取仅预览不入库"],

    ["已存在多条采集配置", "",
     "点击界面右上方的“顺序采集全部”按钮",
     "按配置序号从小到大依次请求全部采集配置，解析预览区汇总显示所有配置采集到的原始数据。",
     "与预期结果一致", ""],

    ["存在步骤5或步骤6的解析预览数据", "",
     "点击解析预览中某条记录下方的“本条采集内容”展开项",
     "展开显示该条采集原文（JSON 文本）。",
     "与预期结果一致", ""],

    ["", ["标签：告警码、", "选中文本：ALM-1001"],
     "在展开的采集原文中选中一段值文本（如告警码 ALM-1001），在“标签”输入框中输入标签名，点击“暂存选中”按钮。",
     "弹出成功提示“已暂存”，标签暂存区新增一条记录，标签、值与输入及选中内容一致，状态显示“待入库”。",
     "与预期结果一致", "选中内容须在当前这条采集原文范围内"],

    ["已存在步骤8暂存的标签记录", ["标签：告警编码、", "值：ALM-1001"],
     "在标签暂存区点击该记录后面的“修改”按钮，在弹出窗口中修改标签与值，点击保存。",
     "标签暂存区列表显示修改后的标签与值，信息与输入一致。",
     "与预期结果一致", ""],

    ["", "",
     "在标签暂存区点击该记录后面的“入库”按钮",
     "弹出成功提示“已入库”，该记录从标签暂存区移除，“已入库数据”区按天聚合新增该条数据，对应日期分组条数加1。",
     "与预期结果一致", "同日同标签将自动覆盖"],

    ["标签暂存区仍存在待入库记录", "",
     "点击标签暂存区右上方的“一键入库”按钮",
     "弹出成功提示“已入库 N 条”，标签暂存区清空，已入库数据按（同日+同标签）规则去重后写入。",
     "与预期结果一致", ""],

    ["存在已入库数据", "",
     "在“已入库数据”区点击对应日期行展开",
     "展开显示该天入库明细，含所属配置、标签、值、来源、类别、入库类别、时间等，与入库数据一致。",
     "与预期结果一致", ""],

    ["", "接口URL：http://192.168.11.87:8080/api/ops/none（不可访问）",
     "新增一条接口URL不可访问的采集配置，点击该配置后面的“采集”按钮。",
     "弹出失败提示，提示“HTTP 请求失败”等错误信息，解析预览区不新增数据。",
     "与预期结果一致", "异常处理"],
]

CAPTION = "表  数据提取"
PREMISE = [
    "前提数据：步骤3采集配置接口（运维告警接口）返回示例如下：",
    '{ "code": 0, "data": { "items": [',
    '  {"id": 1, "alarm_code": "ALM-1001", "level": "严重", "source": "运维管理系统", "title": "服务器CPU过载告警", "occur_time": "2026-05-26 08:30:00"},',
    '  {"id": 2, "alarm_code": "ALM-1002", "level": "一般", "source": "运维管理系统", "title": "磁盘空间不足告警", "occur_time": "2026-05-26 09:15:00"}',
    "] } }",
]


def set_para_text(p, text):
    """在 <w:p> 上保留首个 run 的 rPr，设置其文本，删除多余 run。"""
    runs = p.findall(qn("w:r"))
    if runs:
        first = runs[0]
        for r in runs[1:]:
            p.remove(r)
        for t in first.findall(qn("w:t")):
            first.remove(t)
        for br in first.findall(qn("w:br")):
            first.remove(br)
        t = first.makeelement(qn("w:t"), {})
        t.set(qn("xml:space"), "preserve")
        t.text = text
        first.append(t)
    else:
        r = p.makeelement(qn("w:r"), {})
        t = p.makeelement(qn("w:t"), {})
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p.append(r)


def set_tc(tc, content):
    """设置单元格文本，支持多行（list[str] → 多段落），保留段落/字体格式。"""
    lines = content if isinstance(content, list) else [content]
    if not lines:
        lines = [""]
    ps = tc.findall(qn("w:p"))
    template_p = ps[0]
    for p in ps[1:]:
        tc.remove(p)
    set_para_text(template_p, lines[0])
    prev = template_p
    for ln in lines[1:]:
        newp = copy.deepcopy(template_p)
        set_para_text(newp, ln)
        prev.addnext(newp)
        prev = newp


def main():
    d = docx.Document(SRC)

    # 1) 定位 数据提取 标题段落
    heading = None
    for p in d.paragraphs:
        if "DYSJCJYGL_SJTQ" in p.text:
            heading = p
            break
    if heading is None:
        raise SystemExit("未找到 数据提取 标题段落")

    # 2) 定位模板表 TABLE#13（排班查询用例表）与一个 Normal 标题段落模板
    src_tbl = d.tables[13]
    assert "用例名称" in src_tbl.cell(0, 0).text and "排班查询" in src_tbl.cell(0, 2).text

    caption_template_p = None
    for p in d.paragraphs:
        if p.text.strip().startswith("表 9") and "值班登记" in p.text:
            caption_template_p = p._p
            break
    if caption_template_p is None:
        # 退而取任一 Normal 段落
        for p in d.paragraphs:
            if p.style.name == "Normal" and p.text.strip():
                caption_template_p = p._p
                break

    # 3) 克隆模板表
    new_tbl = copy.deepcopy(src_tbl._tbl)
    trs = new_tbl.findall(qn("w:tr"))

    # 行0：用例名称/标识 → 第二个 tc 写入用例标识
    set_tc(trs[0].findall(qn("w:tc"))[1], CASE_ID)
    # 行1：用例说明 → 第二个 tc 写入说明
    set_tc(trs[1].findall(qn("w:tc"))[1], CASE_DESC)
    # 行2：表头（步骤|前提和约束|输入|目的和动作|预期结果|评估准则|备注）保持不变

    # 取一个步骤行作为模板
    step_template_tr = copy.deepcopy(trs[3])
    # 删除原有步骤行（索引 >=3）
    for tr in trs[3:]:
        new_tbl.remove(tr)

    # 4) 生成步骤行
    for idx, step in enumerate(STEPS, start=1):
        tr = copy.deepcopy(step_template_tr)
        tcs = tr.findall(qn("w:tc"))
        set_tc(tcs[0], str(idx))
        for j in range(6):
            set_tc(tcs[j + 1], step[j])
        new_tbl.append(tr)

    # 5) 构造 caption 与 前提数据 段落
    def make_para(text):
        np = copy.deepcopy(caption_template_p)
        set_para_text(np, text)
        return np

    caption_p = make_para(CAPTION)
    premise_ps = [make_para(t) for t in PREMISE]

    # 6) 依次插入：heading → caption → table → 前提数据段落
    cur = heading._p
    for el in [caption_p, new_tbl] + premise_ps:
        cur.addnext(el)
        cur = el

    import os
    out = SRC
    try:
        d.save(out)
    except PermissionError:
        out = "系统出厂（所）考核大纲模版_已补充.docx"
        d.save(out)
        print("原文件被占用（Word 打开中），已另存为:", out)
    print("OK：已插入数据提取测试用例表，共", len(STEPS), "个步骤")


if __name__ == "__main__":
    main()
